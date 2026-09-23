"""Сборка корпуса data/kb: MD + HTML + PDF + DOCX из страниц музея."""

from __future__ import annotations

import html
import shutil
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from docx import Document as DocxDocument

from app.services.museum_chunks import CATEGORY_BY_STEM

SRC = ROOT / "data" / "museum"
DST = ROOT / "data" / "kb"
DOCX_STEMS = {"hours", "contacts", "tickets", "rules", "hermitage_kazan", "natural_history"}


def _write_html(path: Path, title: str, text: str) -> None:
    body = "<br/>\n".join(html.escape(line) if line else "<br/>" for line in text.splitlines())
    path.write_text(
        f"<!DOCTYPE html><html lang='ru'><head><meta charset='utf-8'><title>{html.escape(title)}</title>"
        f"</head><body><h1>{html.escape(title)}</h1><article>{body}</article></body></html>",
        encoding="utf-8",
    )


def _write_docx(path: Path, title: str, text: str) -> None:
    doc = DocxDocument()
    doc.core_properties.author = "Музей-заповедник Казанский Кремль"
    doc.add_heading(title, level=1)
    for para in text.split("\n\n"):
        if para.strip():
            doc.add_paragraph(para.strip())
    doc.save(path)


def _write_pdf(path: Path, title: str, text: str) -> None:
    import pymupdf

    font_path = Path(r"C:\Windows\Fonts\arial.ttf")
    doc = pymupdf.open()
    page = doc.new_page()
    fontname = "helv"
    if font_path.exists():
        page.insert_font(fontname="f0", fontfile=str(font_path))
        fontname = "f0"
    content = f"{title}\n\n{text}"
    rect = pymupdf.Rect(48, 48, 547, 780)
    page.insert_textbox(rect, content[:4000], fontname=fontname, fontsize=10, align=0)
    doc.save(path, garbage=4, deflate=True)
    doc.close()


def main() -> None:
    if DST.exists():
        shutil.rmtree(DST)
    files = 0
    for src in sorted(SRC.glob("*.md")):
        category = CATEGORY_BY_STEM.get(src.stem, "other")
        folder = DST / category
        folder.mkdir(parents=True, exist_ok=True)
        text = src.read_text(encoding="utf-8")
        title = src.stem.replace("_", " ")
        shutil.copy2(src, folder / src.name)
        files += 1
        _write_html(folder / f"{src.stem}.html", title, text)
        files += 1
        _write_pdf(folder / f"{src.stem}.pdf", title, text)
        files += 1
        if src.stem in DOCX_STEMS:
            _write_docx(folder / f"{src.stem}.docx", title, text)
            files += 1
    (ROOT / "data" / "uploads").mkdir(parents=True, exist_ok=True)
    write_inventory(DST)
    print(f"wrote {files} files under {DST}")


def write_inventory(dst: Path) -> None:
    rows = [p for p in dst.rglob("*") if p.is_file()]
    by_ext = Counter(p.suffix.lower() or "(no ext)" for p in rows)
    total = sum(p.stat().st_size for p in rows)
    lines = [
        "# Инвентаризация корпуса (ДЗ 5.5)",
        "",
        "Источник: `data/museum/*.md` → `data/kb/<category>/` через `python scripts/download_data.py`.",
        "Категория берётся из пути (`hours`, `museums`, `tickets`, …).",
        "",
        f"- **Файлов:** {len(rows)}",
        f"- **Общий размер:** {total} байт ({total / 1024:.1f} КБ)",
        "- **Форматы:**",
    ]
    for ext, count in sorted(by_ext.items()):
        size = sum(p.stat().st_size for p in rows if (p.suffix.lower() or "(no ext)") == ext)
        lines.append(f"  - `{ext}`: {count} файлов, {size} байт")
    lines.extend(
        [
            "",
            "Загрузка посетителем: `POST /documents/upload` кладёт файл в `data/uploads/` и индексирует фоном.",
            "",
            "Повторный ingest: `python scripts/ingest.py data/`.",
            "",
        ]
    )
    out = ROOT / "docs" / "data_inventory.md"
    out.write_text("\n".join(lines), encoding="utf-8")


if __name__ == "__main__":
    main()
